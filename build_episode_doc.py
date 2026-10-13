from pathlib import Path
import json,re,collections,textwrap
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT=Path('E:/v2t/output/S01E01')
data=json.loads((ROOT/'S01E01_分角色台词.json').read_text(encoding='utf-8'))
rows=data['rows']
def stamp(t,srt=False):
    n=round(t*1000);h,n=divmod(n,3600000);m,n=divmod(n,60000);s,ms=divmod(n,1000)
    return f'{h:02}:{m:02}:{s:02},{ms:03}' if srt else f'{m:02}:{s:02}'
def start_stamp(t):
    tenths=round(t*10);minutes,tenths=divmod(tenths,600);seconds,tenth=divmod(tenths,10)
    return f'{minutes:02}:{seconds:02}.{tenth}'
def section(t):
    if t<29:return '片头呼名'
    if t<86.5:return '钢琴与邦戈鼓游戏'
    if t<91:return '片名报幕'
    if t<204:return '魔法木琴与恶作剧'
    if t<257:return '妈妈提醒轮流玩'
    if t<332.5:return '争抢木琴与爸爸反击'
    if t<370:return 'Bingo 说出自己的感受'
    return '姐妹合作与爸爸喷泉'

doc=Document()
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.left_margin=sec.right_margin=Inches(.7)
sec.top_margin=sec.bottom_margin=Inches(.65)
sec.footer_distance=Inches(.28)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
    style=doc.styles[name];style.font.name='Arial';style.font.color.rgb=RGBColor(0,0,0)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Microsoft YaHei')
normal=doc.styles['Normal'];normal.font.size=Pt(11)
normal.paragraph_format.line_spacing=1.08;normal.paragraph_format.space_after=Pt(5)
normal.paragraph_format.widow_control=True
doc.styles['Title'].font.size=Pt(23);doc.styles['Title'].paragraph_format.space_after=Pt(5)
doc.styles['Subtitle'].font.size=Pt(12);doc.styles['Subtitle'].paragraph_format.space_after=Pt(11)
doc.styles['Heading 1'].font.size=Pt(13)
doc.styles['Heading 1'].paragraph_format.space_before=Pt(11)
doc.styles['Heading 1'].paragraph_format.space_after=Pt(6)

doc.add_paragraph('Bluey 第一季第一集分角色台词','Title')
doc.add_paragraph('Magic Xylophone 魔法木琴','Subtitle')
doc.add_paragraph('英文台词按视频时间顺序排列，分别标注角色、片头合声与片名报幕。对白保留原句及口语表达，角色在画外说话时仍记在该角色名下。')
p=doc.add_paragraph('角色：Bluey 布鲁伊；Bingo 宾果；爸爸 Dad（Bandit）；妈妈 Mum（Chilli）。')
p=doc.add_paragraph('素材：S01E01.mp4 与 S01E01.mp3，时长约 07:18。正文时间码标注台词起点，字幕文件保留起止时间；均从视频开头计时，允许半秒左右偏差。')
p=doc.add_paragraph('说明：01:26 的片名报幕单列为“旁白／片名报幕”，不指定配音身份。哼唱和有字幕的拟声已标注；背景音乐、一般笑声及片尾演职员表不作台词。06:03 附近的哼唱有一处歌词待核。')
for p in doc.paragraphs[-2:]:
    for r in p.runs:r.font.size=Pt(9.5)

docrows=[]
for row in rows:
    if (docrows and docrows[-1]['speaker']==row['speaker'] and docrows[-1]['kind']==row['kind']
            and row['start']-docrows[-1]['end']<.8 and section(row['start'])==section(docrows[-1]['start'])
            and not re.search(r'[.!?]$',docrows[-1]['text'])):
        docrows[-1]['text']+=' '+row['text'];docrows[-1]['end']=row['end']
    else:docrows.append(dict(row))
last=None
for row in docrows:
    current=section(row['start'])
    if current!=last:doc.add_paragraph(current,'Heading 1');last=current
    p=doc.add_paragraph()
    p.paragraph_format.keep_together=True
    r=p.add_run(f"{start_stamp(row['start'])}  ");r.font.size=Pt(9)
    label=row['speaker']
    if row['kind'] in ['哼唱','拟声']:label+=f"（{row['note'] if row['kind']=='拟声' else '哼唱'}）"
    r=p.add_run(label+'：');r.bold=True
    r=p.add_run(row['text'])
    if '待核' in row['text']:
        r.font.color.rgb=RGBColor(139,69,19)

footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=footer.add_run('');r.font.size=Pt(9)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
doc.core_properties.title='Bluey 第一季第一集 魔法木琴 分角色台词'
doc.core_properties.subject='英文原文与说话人标注'
doc.core_properties.author='Codex'
for border in list(doc.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
output=ROOT/'Bluey_S01E01_魔法木琴_分角色台词.docx';doc.save(output)

text_lines=['Bluey S01E01 — Magic Xylophone 魔法木琴','英文分角色台词','时间码对应用户提供视频；06:03 附近哼唱有一处歌词待核。','']
srt=[]
for i,row in enumerate(rows,1):
    text_lines.append(f"[{stamp(row['start'])}–{stamp(row['end'])}] {row['speaker']}：{row['text']}")
    subtitle=f"{row['speaker']}：{row['text']}"
    srt.append(f"{i}\n{stamp(row['start'],True)} --> {stamp(row['end'],True)}\n{subtitle}\n")
(ROOT/'Bluey_S01E01_分角色台词.txt').write_text('\n'.join(text_lines)+'\n',encoding='utf-8-sig')
(ROOT/'Bluey_S01E01_分角色字幕.srt').write_text('\n'.join(srt),encoding='utf-8-sig')

assert all(0<=r['start']<r['end']<=data['duration'] for r in rows)
assert all(a['end']<=b['start']+.001 for a,b in zip(rows,rows[1:]))
assert sum(r['kind']=='旁白／报幕' for r in rows)==1
assert len([r for r in rows if '待核' in r['text']])==1
included={i for r in rows for i in r['source_caption_ids']}
assert set(range(1,135))-included=={54,124}
counts=collections.Counter(r['speaker'] for r in rows)
report={'transcript_entries':len(rows),'counts':dict(counts),'uncertainties':data['uncertainties'],'coverage':'All 134 pre-credit OCR groups handled, excluding false positive #54 and merging duplicate #124 into #123; one short caption and one uncaptioned sung line added.'}
(ROOT/'verification_report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(report,ensure_ascii=True))
print(str(output).encode('ascii','backslashreplace').decode())
