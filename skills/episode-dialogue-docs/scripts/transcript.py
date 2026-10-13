"""Validate reviewed dialogue data and export consistent DOCX/TXT/SRT/CSV files."""
import argparse
import csv
import hashlib
import json
import math
from pathlib import Path
import re
import sys

DELIVERIES = {
    'dialogue': '', 'offscreen_dialogue': '画外对白', 'narration': '旁白',
    'title_announcement': '片名报幕', 'chorus': '合声',
    'singing': '哼唱', 'vocalization': '拟声',
}


def read_json(path):
    return json.loads(Path(path).read_text(encoding='utf-8-sig'))


def available(path):
    path = Path(path)
    if path.exists():
        raise FileExistsError(f'Refusing to overwrite: {path}')
    path.parent.mkdir(parents=True, exist_ok=True)
    return path


def save_json(path, value):
    with available(path).open('x', encoding='utf-8') as handle:
        json.dump(value, handle, ensure_ascii=False, indent=2, allow_nan=False)


def nonempty(value):
    return isinstance(value, str) and bool(value.strip())


def finite(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value)


def registry_errors(registry):
    errors = []
    if registry.get('schema_version') != 1 or not nonempty(registry.get('series_id')):
        errors.append('Invalid registry schema_version or series_id')
    if not isinstance(registry.get('revision'), int) or registry['revision'] < 1:
        errors.append('Registry revision must be a positive integer')
    seen = set()
    characters = registry.get('characters', [])
    if not isinstance(characters, list):
        return errors + ['characters must be a list']
    for character in characters:
        cid = character.get('id', '')
        if not re.fullmatch(r'[a-z0-9][a-z0-9_-]*', cid) or cid in seen:
            errors.append(f'Invalid/duplicate character ID: {cid}')
        seen.add(cid)
        if not nonempty(character.get('display_name')):
            errors.append(f'{cid}: display_name required')
        if character.get('status') not in {'confirmed', 'provisional'}:
            errors.append(f'{cid}: invalid status')
        refs = character.get('references', [])
        if character.get('status') == 'confirmed' and not refs:
            errors.append(f'{cid}: confirmed character needs a reference')
        for ref in refs:
            if not all(nonempty(ref.get(k)) for k in ['source', 'episode', 'basis', 'language', 'edition']):
                errors.append(f'{cid}: incomplete reference')
            if not re.fullmatch(r'[0-9a-f]{64}', ref.get('sha256', '')):
                errors.append(f'{cid}: reference requires SHA256')
            if not (finite(ref.get('start')) and finite(ref.get('end')) and 0 <= ref['start'] < ref['end']):
                errors.append(f'{cid}: invalid reference time')
    return errors


def validate(data, registry):
    errors = registry_errors(registry)
    characters = {c['id']: c for c in registry.get('characters', [])}
    if data.get('schema_version') != 1 or data.get('series_id') != registry.get('series_id'):
        errors.append('Transcript schema/series does not match registry')
    for field in ['episode', 'title', 'language', 'timecode_note', 'scope_note']:
        if not nonempty(data.get(field)):
            errors.append(f'Missing {field}')
    duration = data.get('duration')
    if not finite(duration) or duration <= 0:
        errors.append('Invalid duration')
        duration = 0
    if not data.get('sources'):
        errors.append('Missing media sources')
    for source in data.get('sources', []):
        if not nonempty(source.get('path')) or not re.fullmatch(r'[0-9a-f]{64}', source.get('sha256', '')):
            errors.append('Source requires path and SHA256')
    segments = data.get('segments', [])
    if not segments:
        errors.append('Transcript has no segments')
    used = set()
    ids = set()
    last_start, latest_end = -1.0, -1.0
    uncertain = []
    for row in segments:
        rid = row.get('id')
        prefix = f'Segment {rid}: '
        if not isinstance(rid, int) or rid < 1 or rid in ids:
            errors.append(prefix + 'invalid/duplicate ID')
        ids.add(rid)
        start, end = row.get('start'), row.get('end')
        if not (finite(start) and finite(end) and 0 <= start < end <= duration):
            errors.append(prefix + 'invalid time range')
        else:
            if round(end * 1000) <= round(start * 1000):
                errors.append(prefix + 'time range collapses at SRT precision')
            if start < last_start:
                errors.append(prefix + 'not sorted by start time')
            if start < latest_end - .001 and not nonempty(row.get('overlap_note')):
                errors.append(prefix + 'unexplained overlap')
            last_start, latest_end = start, max(latest_end, end)
        speakers = row.get('speaker_ids')
        if not isinstance(speakers, list) or not all(isinstance(s, str) for s in speakers):
            errors.append(prefix + 'speaker_ids must be a list of IDs')
            speakers = []
        if len(speakers) != len(set(speakers)):
            errors.append(prefix + 'duplicate speaker ID')
        for speaker in speakers:
            if speaker not in characters:
                errors.append(prefix + f'unknown character {speaker}')
            elif characters[speaker]['status'] == 'provisional' and row.get('speaker_status') == 'reviewed':
                errors.append(prefix + 'provisional identity cannot be marked reviewed')
        if row.get('delivery') not in DELIVERIES:
            errors.append(prefix + 'invalid delivery')
        ts, ss = row.get('text_status'), row.get('speaker_status')
        if ts not in {'reviewed', 'uncertain'} or ss not in {'reviewed', 'uncertain', 'group'}:
            errors.append(prefix + 'invalid review status')
        if not speakers and ss == 'reviewed':
            errors.append(prefix + 'unknown speaker cannot be marked reviewed')
        if ss == 'group' and (speakers or row.get('delivery') != 'chorus' or not nonempty(row.get('group_label'))):
            errors.append(prefix + 'unidentified group needs chorus and group_label without character IDs')
        if not nonempty(row.get('text')):
            errors.append(prefix + 'empty dialogue text')
        elif any(token in row['text'] for token in ['待核', '听不清']) and ts != 'uncertain':
            errors.append(prefix + 'unclear text cannot be marked reviewed')
        evidence = row.get('evidence')
        if not isinstance(evidence, list) or not evidence or not all(nonempty(e) for e in evidence):
            errors.append(prefix + 'review evidence required')
        if ts == 'uncertain' or ss == 'uncertain':
            if not nonempty(row.get('review_note')):
                errors.append(prefix + 'uncertainty needs a review_note')
            uncertain.append({'id': rid, 'start': start, 'text_status': ts, 'speaker_status': ss,
                              'reason': row.get('review_note', '')})
        refs = row.get('source_ids')
        if not isinstance(refs, list) or not all(nonempty(ref) for ref in refs):
            errors.append(prefix + 'source_ids must be a list')
        else:
            used.update(refs)
    coverage = data.get('coverage', {})
    expected = coverage.get('source_ids')
    if not isinstance(expected, list) or not all(nonempty(x) for x in expected):
        errors.append('coverage.source_ids must be a list')
        expected = []
    if len(expected) != len(set(expected)):
        errors.append('Duplicate coverage ID')
    excluded = set()
    for item in coverage.get('excluded', []):
        key = item.get('id')
        if not nonempty(key) or not nonempty(item.get('reason')) or key in excluded:
            errors.append('Invalid/duplicate excluded source or missing reason')
        excluded.add(key)
    if used & excluded:
        errors.append('Source ID both used and excluded')
    if set(expected) != used | excluded:
        errors.append(f'Coverage mismatch: missing={sorted(set(expected)-(used|excluded))}, extra={sorted((used|excluded)-set(expected))}')
    if not nonempty(coverage.get('review_note')):
        errors.append('Coverage review_note required')
    return {'ok': not errors, 'errors': errors, 'entries': len(segments), 'uncertainties': uncertain,
            'registry_revision': registry.get('revision'), 'visual_review': 'not_performed_by_validator'}


def timestamp(seconds, srt=False):
    total = round(seconds * 1000)
    hours, total = divmod(total, 3600000)
    minutes, total = divmod(total, 60000)
    secs, millis = divmod(total, 1000)
    if srt:
        return f'{hours:02}:{minutes:02}:{secs:02},{millis:03}'
    return f'{hours:02}:{minutes:02}:{secs:02}.{millis:03}'


def label(row, registry):
    names = {c['id']: c['display_name'] for c in registry['characters']}
    result = ' 和 '.join(names[s] for s in row['speaker_ids'])
    if row['speaker_status'] == 'uncertain':
        result = (result + '（身份待核）') if result else '说话人待核'
    if row['speaker_status'] == 'group':
        result = row['group_label']
    delivery = DELIVERIES[row['delivery']]
    if delivery and delivery not in result:
        result += f'（{delivery}）'
    return result


def dialogue(row):
    value = row['text']
    if row['text_status'] == 'uncertain' and not any(t in value for t in ['待核', '听不清']):
        value += ' [文字待核]'
    return value


def build_docx(data, registry, output):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Inches(.7)
    sec.top_margin = sec.bottom_margin = Inches(.65)
    for name in ['Normal', 'Title', 'Subtitle', 'Heading 1']:
        style = doc.styles[name]
        style.font.name = 'Arial'
        style.font.color.rgb = RGBColor(0, 0, 0)
        fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        for attr, font in [('ascii', 'Arial'), ('hAnsi', 'Arial'), ('eastAsia', 'Microsoft YaHei')]:
            fonts.set(qn('w:' + attr), font)
    normal = doc.styles['Normal']
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.space_after = Pt(5)
    doc.styles['Title'].font.size = Pt(23)
    doc.styles['Subtitle'].font.size = Pt(12)
    doc.styles['Heading 1'].font.size = Pt(13)
    heading = lambda text: ' '.join(re.sub(r'[^\w\s]', ' ', text).split())
    doc.add_paragraph(heading(f"{registry.get('series_name', data['series_id'])} {data['episode']} 分角色台词"), 'Title')
    doc.add_paragraph(heading(data['title'] + ' ' + data.get('title_zh', '')), 'Subtitle')
    doc.add_paragraph('按原片时间顺序排列，分别标注说话人和发声形式。角色在画外说话时仍归入该角色。')
    used = {s for row in data['segments'] for s in row['speaker_ids']}
    roster = []
    for character in registry['characters']:
        if character['id'] in used:
            value, chinese = character['display_name'], character.get('name_zh', '')
            roster.append(value if not chinese or chinese in value else value + ' ' + chinese)
    doc.add_paragraph('角色：' + '；'.join(roster) + '。')
    doc.add_paragraph(data['timecode_note'])
    doc.add_paragraph(data['scope_note'])
    previous = None
    for row in data['segments']:
        section = row.get('section')
        if section and section != previous:
            doc.add_paragraph(heading(section), 'Heading 1')
        previous = section
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.keep_together = True
        paragraph.add_run(timestamp(row['start']) + '  ').font.size = Pt(9)
        paragraph.add_run(label(row, registry) + '：').bold = True
        run = paragraph.add_run(dialogue(row))
        if row['text_status'] == 'uncertain' or row['speaker_status'] == 'uncertain':
            run.font.color.rgb = RGBColor(139, 69, 19)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    for root in [doc.styles.element, doc.element]:
        for border in list(root.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    doc.core_properties.title = f"{data['episode']} {data['title']} 分角色台词"
    doc.core_properties.author = 'Codex'
    doc.save(output)


def export(data, registry, directory, formats):
    report = validate(data, registry)
    if not report['ok']:
        raise ValueError('\n'.join(report['errors']))
    directory = Path(directory)
    stem = re.sub(r'[^\w.-]+', '_', f"{data['series_id']}_{data['episode']}_分角色台词")
    targets = {fmt: directory / f'{stem}.{fmt}' for fmt in formats}
    report_path = directory / 'validation.json'
    for target in [*targets.values(), report_path]:
        available(target)
    lines = [f"{data['episode']} {data['title']}", data['timecode_note'], data['scope_note'], '']
    cues = []
    for i, row in enumerate(data['segments'], 1):
        body = f'{label(row, registry)}：{dialogue(row)}'
        lines.append(f"[{timestamp(row['start'])}–{timestamp(row['end'])}] {body}")
        cues.append(f"{i}\n{timestamp(row['start'], True)} --> {timestamp(row['end'], True)}\n{body}\n")
    if 'txt' in targets:
        targets['txt'].write_text('\n'.join(lines) + '\n', encoding='utf-8-sig')
    if 'srt' in targets:
        targets['srt'].write_text('\n'.join(cues), encoding='utf-8-sig')
    if 'docx' in targets:
        build_docx(data, registry, targets['docx'])
    report['output_files'] = {key: str(value) for key, value in targets.items()}
    report['delivery_state'] = 'awaiting_content_and_visual_review'
    save_json(report_path, report)
    return report


def registry_csv(registry, output):
    errors = registry_errors(registry)
    if errors:
        raise ValueError('\n'.join(errors))
    with available(output).open('x', encoding='utf-8-sig', newline='') as handle:
        writer = csv.writer(handle)
        writer.writerow(['角色编号', '英文名', '中文名', '文档显示名', '别名', '状态', '外观参考', '声音备注', '参考片段'])
        for char in registry['characters']:
            refs = '; '.join(f"{r['episode']} {r['start']:.3f}-{r['end']:.3f}s {r['source']}" for r in char.get('references', []))
            values = [char['id'], char.get('name_en', ''), char.get('name_zh', ''), char['display_name'],
                      ' / '.join(char.get('aliases', [])), char['status'], char.get('visual_notes', ''), char.get('voice_notes', ''), refs]
            writer.writerow(["'" + value if value.startswith(('=', '+', '-', '@')) else value for value in values])


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    commands = parser.add_subparsers(dest='command', required=True)
    for command in ['validate', 'export', 'registry-csv']:
        sub = commands.add_parser(command)
        sub.add_argument('--registry', required=True)
        if command != 'registry-csv':
            sub.add_argument('--transcript', required=True)
        if command == 'validate':
            sub.add_argument('--report')
        elif command == 'export':
            sub.add_argument('--out-dir', required=True)
            sub.add_argument('--formats', nargs='+', choices=['docx', 'txt', 'srt'], default=['docx', 'txt', 'srt'])
        else:
            sub.add_argument('--output', required=True)
    args = parser.parse_args()
    registry = read_json(args.registry)
    if args.command == 'registry-csv':
        registry_csv(registry, args.output)
        print(args.output)
        return
    data = read_json(args.transcript)
    if args.command == 'export':
        result = export(data, registry, args.out_dir, args.formats)
    else:
        result = validate(data, registry)
        result['registry_sha256'] = hashlib.sha256(Path(args.registry).read_bytes()).hexdigest()
        result['transcript_sha256'] = hashlib.sha256(Path(args.transcript).read_bytes()).hexdigest()
        if args.report:
            save_json(args.report, result)
    print(json.dumps(result, ensure_ascii=True, indent=2))
    if not result['ok']:
        sys.exit(1)


if __name__ == '__main__':
    main()
