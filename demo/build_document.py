from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parent
data = json.loads((ROOT / 'transcript_final.json').read_text(encoding='utf-8'))

def stamp(seconds, srt=False):
    n = round(seconds * 1000)
    h, n = divmod(n, 3600000)
    m, n = divmod(n, 60000)
    s, ms = divmod(n, 1000)
    return f'{h:02}:{m:02}:{s:02},{ms:03}' if srt else f'{m:02}:{s:02}'

doc = Document()
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin, sec.bottom_margin = Inches(.7), Inches(.7)
sec.left_margin, sec.right_margin = Inches(.8), Inches(.8)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2']:
    style = doc.styles[name]
    style.font.name = 'Calibri'
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Microsoft YaHei')
normal = doc.styles['Normal']
normal.font.size = Pt(11)
normal.paragraph_format.line_spacing = 1.12
normal.paragraph_format.space_after = Pt(7)
doc.styles['Title'].font.size = Pt(23)
doc.styles['Title'].paragraph_format.space_after = Pt(9)
doc.styles['Heading 1'].font.size = Pt(13)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(13)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(7)

doc.add_paragraph('Window on the World 英文台词', 'Title')
doc.add_paragraph('NASA 宇航员 Scott Kelly 介绍国际空间站观景舱。这份台词稿覆盖整段视频中的讲话，时间码从视频开始计算，可用于对照原片。')
p = doc.add_paragraph('视频时长 03:19  |  语言 英语  |  主讲人 Scott Kelly')
p.paragraph_format.space_after = Pt(4)

def link(paragraph, label, url):
    rel = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    el = OxmlElement('w:hyperlink')
    el.set(qn('r:id'), rel)
    run = OxmlElement('w:r')
    props = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), '0563C1'); props.append(color)
    underline = OxmlElement('w:u'); underline.set(qn('w:val'), 'single'); props.append(underline)
    run.append(props)
    txt = OxmlElement('w:t'); txt.text = label; run.append(txt)
    el.append(run); paragraph._p.append(el)

p = doc.add_paragraph('原片来源 ')
link(p, 'NASA Johnson 原视频', 'https://www.youtube.com/watch?v=U54RMrBK9vU')
p.add_run('  |  ')
link(p, '可播放的视频与来源说明', 'https://commons.wikimedia.org/wiki/File:Window_on_the_World.webm')
p = doc.add_paragraph('转写说明：由视频音轨自动识别，整理标点与分段。时间码为自动估计，可能略有偏差；准确措辞以原音频为准。')
for run in p.runs:
    run.font.size = Pt(9)
doc.add_paragraph('带时间码的英文全文', 'Heading 1')

for group in data['paragraphs']:
    p = doc.add_paragraph()
    lead = p.add_run(f"[{stamp(group['start'])} - {stamp(group['end'])}]  ")
    lead.bold = True
    lead.font.size = Pt(10)
    p.add_run(group['text'])
    p.paragraph_format.widow_control = True

doc.core_properties.title = 'Window on the World 英文台词'
doc.core_properties.subject = 'NASA 英文视频转写演示'
doc.core_properties.author = 'Codex'
for border in list(doc.element.iter(qn('w:pBdr'))):
    border.getparent().remove(border)
output = ROOT / 'Window_on_the_World_英文台词.docx'
doc.save(output)

srt = []
for i, row in enumerate(data['segments'], 1):
    srt.append(f"{i}\n{stamp(row['start'], True)} --> {stamp(row['end'], True)}\n{row['text']}\n")
(ROOT / 'Window_on_the_World.en.srt').write_text('\n'.join(srt), encoding='utf-8')
print(output)
print(f"Paragraphs: {len(data['paragraphs'])}; subtitle entries: {len(data['segments'])}")
